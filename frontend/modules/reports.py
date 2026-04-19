import os
import datetime
from PyQt6.QtWidgets import (QVBoxLayout, QHBoxLayout, QTableWidget, QTableWidgetItem,
                             QHeaderView, QLabel, QWidget, QTabWidget, QPushButton,
                             QFileDialog, QMessageBox, QFrame, QSplitter)
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QColor, QBrush, QFont
from modules.base import BaseModule, parse_brand_model
from api_client import api, APIError

SEASON_COLORS = {"Kislik": "#74c0fc", "Yazlik": "#ffd43b", "4 Mevsim": "#8ce99a"}


def _now_str():
    return datetime.datetime.now().strftime("%Y%m%d_%H%M%S")

def _right(text):
    item = QTableWidgetItem(text)
    item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
    return item

def _center(text):
    item = QTableWidgetItem(text)
    item.setTextAlignment(Qt.AlignmentFlag.AlignCenter | Qt.AlignmentFlag.AlignVCenter)
    return item

def _stat_card(label, value, color="#cdd6f4"):
    w = QFrame()
    w.setStyleSheet(f"background:#313244; border-radius:6px; padding:4px;")
    lay = QVBoxLayout(w)
    lay.setContentsMargins(10, 6, 10, 6)
    lbl = QLabel(label)
    lbl.setStyleSheet("font-size:10px; color:#a6adc8;")
    val = QLabel(value)
    val.setStyleSheet(f"font-size:15px; font-weight:bold; color:{color};")
    val.setAlignment(Qt.AlignmentFlag.AlignCenter)
    lay.addWidget(lbl)
    lay.addWidget(val)
    return w, val


# ─── Export yardımcıları ──────────────────────────────────────────────────────

def _export_txt(filepath, period_label, data):
    sales = data.get("sales", [])
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(f"{'='*75}\n  {period_label.upper()} SATIS RAPORU\n")
        f.write(f"  {datetime.datetime.now().strftime('%d.%m.%Y %H:%M')}\n{'='*75}\n\n")
        f.write(f"  Toplam Satis   : {data.get('total_sales',0):>12,.2f} TL\n")
        f.write(f"  Toplam Maliyet : {data.get('total_cost',0):>12,.2f} TL\n")
        f.write(f"  Brut Kar       : {data.get('gross_profit',0):>12,.2f} TL\n")
        f.write(f"  Giderler       : {data.get('total_expenses',0):>12,.2f} TL\n")
        f.write(f"  Net Kar        : {data.get('net_profit',0):>12,.2f} TL\n")
        f.write(f"  Zarali Satis   : {data.get('loss_count',0)} adet\n\n")
        f.write(f"{'ID':<8}{'Tarih':<20}{'Odeme':<14}{'Tutar':>10}{'Maliyet':>10}{'Kar/Zarar':>12}\n")
        f.write(f"{'-'*75}\n")
        for s in sales:
            flag = " ⚠" if s.get("is_loss") else ""
            f.write(f"{s['id']:<8}{s['timestamp'][:19].replace('T',' '):<20}"
                    f"{s.get('payment_method','-'):<14}"
                    f"{s['total_amount']:>10,.2f}{s.get('total_cost',0):>10,.2f}"
                    f"{s.get('profit',0):>12,.2f}{flag}\n")
        f.write(f"{'-'*75}\n")
        f.write(f"{'TOPLAM':<42}{data.get('total_sales',0):>10,.2f}"
                f"{data.get('total_cost',0):>10,.2f}{data.get('gross_profit',0):>12,.2f} TL\n")


def _export_excel(filepath, period_label, data):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise RuntimeError("openpyxl kurulu degil: pip install openpyxl")

    sales = data.get("sales", [])
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = period_label
    h_font = Font(bold=True, color="FFFFFF"); h_fill = PatternFill("solid", fgColor="1a56b0")
    t_font = Font(bold=True); t_fill = PatternFill("solid", fgColor="dbeafe")
    loss_fill = PatternFill("solid", fgColor="fee2e2")
    thin = Side(style="thin", color="cccccc"); border = Border(left=thin,right=thin,top=thin,bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A1:F1"); ws["A1"] = f"{period_label} — {datetime.datetime.now().strftime('%d.%m.%Y %H:%M')}"
    ws["A1"].font = Font(bold=True, size=13); ws["A1"].alignment = center; ws.row_dimensions[1].height = 28

    summary = [("Toplam Satis", data.get("total_sales",0)), ("Maliyet", data.get("total_cost",0)),
               ("Brut Kar", data.get("gross_profit",0)), ("Giderler", data.get("total_expenses",0)),
               ("Net Kar", data.get("net_profit",0)), ("Zarali Satis", data.get("loss_count",0))]
    for i, (lbl, val) in enumerate(summary, 2):
        ws.cell(row=i, column=1, value=lbl).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=round(val,2) if isinstance(val,float) else val)
        c.number_format = "#,##0.00"
        if lbl in ("Brut Kar","Net Kar") and isinstance(val,float) and val<0:
            c.font = Font(color="cc0000", bold=True)

    hr = len(summary)+3
    for col, h in enumerate(["ID","Tarih","Odeme","Tutar","Maliyet","Kar/Zarar"],1):
        c = ws.cell(row=hr, column=col, value=h)
        c.font=h_font; c.fill=h_fill; c.alignment=center; c.border=border

    for ri, s in enumerate(sales, hr+1):
        profit = s.get("profit",0.0); is_loss = s.get("is_loss",False)
        rf = loss_fill if is_loss else None
        for col, val in enumerate([s["id"], s["timestamp"][:19].replace("T"," "),
                                    s.get("payment_method","-"),
                                    round(s["total_amount"],2), round(s.get("total_cost",0),2),
                                    round(profit,2)], 1):
            c = ws.cell(row=ri, column=col, value=val); c.border=border
            if rf: c.fill=rf
            if col>=4: c.number_format="#,##0.00"; c.alignment=Alignment(horizontal="right")
            if col==6 and profit<0: c.font=Font(color="cc0000")

    tr = hr+len(sales)+1
    ws.merge_cells(f"A{tr}:C{tr}")
    lbl=ws.cell(row=tr,column=1,value="TOPLAM"); lbl.font=t_font; lbl.fill=t_fill; lbl.alignment=center; lbl.border=border
    for col,val in [(4,data.get("total_sales",0)),(5,data.get("total_cost",0)),(6,data.get("gross_profit",0))]:
        c=ws.cell(row=tr,column=col,value=round(val,2)); c.font=t_font; c.fill=t_fill
        c.number_format="#,##0.00"; c.alignment=Alignment(horizontal="right"); c.border=border
    for ch in "ABCDEF": ws.column_dimensions[ch].width=16
    wb.save(filepath)


def _export_pdf(filepath, period_label, data):
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
    except ImportError:
        raise RuntimeError("reportlab kurulu degil: pip install reportlab")

    sales = data.get("sales",[])
    doc = SimpleDocTemplate(filepath, pagesize=landscape(A4),
                            leftMargin=2*cm,rightMargin=2*cm,topMargin=2*cm,bottomMargin=2*cm)
    styles=getSampleStyleSheet(); story=[]
    story.append(Paragraph(f"{period_label} Satis Raporu", styles["Title"]))
    story.append(Paragraph(datetime.datetime.now().strftime('%d.%m.%Y %H:%M'),
                           ParagraphStyle("s",parent=styles["Normal"],textColor=colors.grey,spaceAfter=10)))
    sd=[["Toplam Satis",f"{data.get('total_sales',0):,.2f} TL",
         "Brut Kar",f"{data.get('gross_profit',0):,.2f} TL",
         "Net Kar",f"{data.get('net_profit',0):,.2f} TL"]]
    st=Table(sd,colWidths=[4*cm]*6)
    st.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1a56b0")),
                             ("TEXTCOLOR",(0,0),(-1,0),colors.white),
                             ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
                             ("ALIGN",(0,0),(-1,-1),"CENTER"),
                             ("GRID",(0,0),(-1,-1),0.5,colors.lightgrey)]))
    story.append(st); story.append(Spacer(1,0.5*cm))
    td=[["ID","Tarih","Odeme","Tutar (TL)","Maliyet (TL)","Kar/Zarar (TL)"]]
    for s in sales:
        td.append([str(s["id"]),s["timestamp"][:19].replace("T"," "),s.get("payment_method","-"),
                   f"{s['total_amount']:,.2f}",f"{s.get('total_cost',0):,.2f}",f"{s.get('profit',0):,.2f}"])
    td.append(["","","TOPLAM",f"{data.get('total_sales',0):,.2f}",
               f"{data.get('total_cost',0):,.2f}",f"{data.get('gross_profit',0):,.2f}"])
    loss_rows=[i+1 for i,s in enumerate(sales) if s.get("is_loss")]
    t=Table(td,colWidths=[2*cm,5*cm,4*cm,4*cm,4*cm,4*cm])
    cmds=[("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1a56b0")),("TEXTCOLOR",(0,0),(-1,0),colors.white),
          ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("ALIGN",(0,0),(-1,-1),"CENTER"),
          ("ALIGN",(3,1),(5,-1),"RIGHT"),("BACKGROUND",(0,-1),(-1,-1),colors.HexColor("#dbeafe")),
          ("FONTNAME",(0,-1),(-1,-1),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),0.5,colors.lightgrey),
          ("ROWBACKGROUNDS",(0,1),(-1,-2),[colors.white,colors.HexColor("#f8faff")])]
    for lr in loss_rows:
        cmds.append(("BACKGROUND",(0,lr),(-1,lr),colors.HexColor("#fee2e2")))
        cmds.append(("TEXTCOLOR",(5,lr),(5,lr),colors.HexColor("#cc0000")))
    t.setStyle(TableStyle(cmds)); story.append(t); doc.build(story)


# ─── Rapor sekmesi widget'ı ───────────────────────────────────────────────────

class PeriodReportTab(QWidget):
    def __init__(self, period_label, api_method):
        super().__init__()
        self.period_label = period_label
        self.api_method   = api_method
        self._data        = {}
        layout = QVBoxLayout(self)

        # Özet kartları
        cards_row = QHBoxLayout()
        self._cards = {}
        card_defs = [
            ("total_sales",    "Toplam Satis",  "#89b4fa"),
            ("total_cost",     "Toplam Maliyet","#a6adc8"),
            ("gross_profit",   "Brut Kar",      "#a6e3a1"),
            ("total_expenses", "Giderler",      "#fab387"),
            ("net_profit",     "Net Kar",       "#a6e3a1"),
            ("loss_count",     "Zarali Satis",  "#f38ba8"),
        ]
        for key, lbl, color in card_defs:
            w, val_lbl = _stat_card(lbl, "—", color)
            self._cards[key] = (w, val_lbl, color)
            cards_row.addWidget(w)
        layout.addLayout(cards_row)

        # Export butonları
        btn_row = QHBoxLayout()
        btn_row.addStretch()
        for label, color, fmt in [("TXT ↓","#6b7280","txt"),("Excel ↓","#166534","xlsx"),("PDF ↓","#991b1b","pdf")]:
            btn = QPushButton(label)
            btn.setStyleSheet(f"background-color:{color};color:white;padding:6px 14px;font-weight:bold;border-radius:4px;")
            btn.clicked.connect(lambda _,f=fmt: self._export(f))
            btn_row.addWidget(btn)
        layout.addLayout(btn_row)

        # Tablo
        self.table = QTableWidget()
        self.table.setColumnCount(6)
        self.table.setHorizontalHeaderLabels(
            ["Satis ID","Tarih","Odeme Turu","Tutar (TL)","Maliyet (TL)","Kar/Zarar (TL)"])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        layout.addWidget(self.table)

    def load(self):
        try:
            data = self.api_method()
            self._data = data
            sales = data.get("sales", [])

            card_vals = {
                "total_sales":    f"{data.get('total_sales',0):,.2f} TL",
                "total_cost":     f"{data.get('total_cost',0):,.2f} TL",
                "gross_profit":   f"{data.get('gross_profit',0):,.2f} TL",
                "total_expenses": f"{data.get('total_expenses',0):,.2f} TL",
                "net_profit":     f"{data.get('net_profit',0):,.2f} TL",
                "loss_count":     f"{data.get('loss_count',0)} adet",
            }
            for key, (w, val_lbl, default_color) in self._cards.items():
                val_lbl.setText(card_vals[key])
                if key in ("gross_profit","net_profit"):
                    v = data.get(key,0)
                    val_lbl.setStyleSheet(
                        f"font-size:15px;font-weight:bold;color:{'#a6e3a1' if v>=0 else '#f38ba8'};")

            self.table.setRowCount(0)
            for row, s in enumerate(sales):
                self.table.insertRow(row)
                profit  = s.get("profit",0.0)
                is_loss = s.get("is_loss",False)
                self.table.setItem(row,0,QTableWidgetItem(str(s["id"])))
                self.table.setItem(row,1,QTableWidgetItem(s["timestamp"][:19].replace("T"," ")))
                self.table.setItem(row,2,QTableWidgetItem(s.get("payment_method","-")))
                self.table.setItem(row,3,_right(f"{s['total_amount']:,.2f}"))
                self.table.setItem(row,4,_right(f"{s.get('total_cost',0):,.2f}"))
                p_item = _right(f"{profit:,.2f}")
                if is_loss:
                    for col in range(6):
                        item = self.table.item(row,col)
                        if item: item.setBackground(QBrush(QColor("#2d1a1a")))
                    p_item.setForeground(QBrush(QColor("#f38ba8")))
                else:
                    p_item.setForeground(QBrush(QColor("#a6e3a1")))
                self.table.setItem(row,5,p_item)

        except APIError as e:
            QMessageBox.critical(self, "Hata", str(e))

    def _export(self, fmt):
        if not self._data.get("sales"):
            QMessageBox.information(self,"Bilgi","Veri yok. Once raporu yukleyin."); return
        ext_map={"txt":"TXT (*.txt)","xlsx":"Excel (*.xlsx)","pdf":"PDF (*.pdf)"}
        filepath,_=QFileDialog.getSaveFileName(self,"Raporu Kaydet",
            f"Rapor_{self.period_label}_{_now_str()}.{fmt}",ext_map[fmt])
        if not filepath: return
        try:
            if fmt=="txt":  _export_txt(filepath,self.period_label,self._data)
            elif fmt=="xlsx": _export_excel(filepath,self.period_label,self._data)
            elif fmt=="pdf":  _export_pdf(filepath,self.period_label,self._data)
            QMessageBox.information(self,"Basarili",f"Kaydedildi:\n{filepath}")
        except Exception as e:
            QMessageBox.critical(self,"Hata",str(e))


class TopProductsTab(QWidget):
    def __init__(self):
        super().__init__()
        layout = QVBoxLayout(self)

        hdr = QLabel("En Cok Satan ve En Karli Urunler")
        hdr.setStyleSheet("font-size:14px;font-weight:bold;margin-bottom:4px;")
        layout.addWidget(hdr)

        self.table = QTableWidget()
        self.table.setColumnCount(9)
        self.table.setHorizontalHeaderLabels(
            ["Urun Adi","Marka","Model","Mevsim","Toplam Adet","Toplam Satis","Toplam Maliyet","Toplam Kar","Durum"])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        layout.addWidget(self.table)

    def load(self):
        try:
            products = api.get_top_products()
            self.table.setRowCount(0)
            for row, p in enumerate(products):
                self.table.insertRow(row)
                is_loss = p.get("is_loss",False)
                profit  = p.get("total_profit",0.0)
                season, model = parse_brand_model(p.get("brand_model") or "")
                season_item = _center(season or "-")
                if season and season in SEASON_COLORS:
                    season_item.setForeground(QBrush(QColor(SEASON_COLORS[season])))
                cells = [
                    QTableWidgetItem(p["name"]),
                    QTableWidgetItem(p.get("brand_name") or "-"),
                    QTableWidgetItem(model or "-"),
                    season_item,
                    _right(str(p["total_qty"])),
                    _right(f"{p['total_revenue']:,.2f} TL"),
                    _right(f"{p['total_cost']:,.2f} TL"),
                    _right(f"{profit:,.2f} TL"),
                    _center("ZARAR" if is_loss else "Karli"),
                ]
                for col, item in enumerate(cells):
                    self.table.setItem(row, col, item)
                if is_loss:
                    for col in range(9):
                        it = self.table.item(row,col)
                        if it: it.setBackground(QBrush(QColor("#2d1a1a")))
                    self.table.item(row,7).setForeground(QBrush(QColor("#f38ba8")))
                    self.table.item(row,8).setForeground(QBrush(QColor("#f38ba8")))
                else:
                    self.table.item(row,7).setForeground(QBrush(QColor("#a6e3a1")))
                    self.table.item(row,8).setForeground(QBrush(QColor("#a6e3a1")))
        except APIError as e:
            QMessageBox.critical(self,"Hata",str(e))


class SupplierSummaryTab(QWidget):
    def __init__(self):
        super().__init__()
        layout = QVBoxLayout(self)
        hdr = QLabel("Tedarkici Bazinda Alim Ozeti")
        hdr.setStyleSheet("font-size:14px;font-weight:bold;margin-bottom:4px;")
        layout.addWidget(hdr)

        self.table = QTableWidget()
        self.table.setColumnCount(4)
        self.table.setHorizontalHeaderLabels(
            ["Tedarikci","Alim Sayisi","Toplam Alim Tutari","Mevcut Borc"])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        layout.addWidget(self.table)

    def load(self):
        try:
            data = api.get_supplier_summary()
            self.table.setRowCount(0)
            for row, s in enumerate(data):
                self.table.insertRow(row)
                debt = s.get("current_debt",0)
                cells = [
                    QTableWidgetItem(s["name"]),
                    _center(str(s["purchase_count"])),
                    _right(f"{s['total_purchases']:,.2f} TL"),
                    _right(f"{debt:,.2f} TL"),
                ]
                for col, item in enumerate(cells):
                    self.table.setItem(row,col,item)
                if debt > 0:
                    self.table.item(row,3).setForeground(QBrush(QColor("#f38ba8")))
        except APIError as e:
            QMessageBox.critical(self,"Hata",str(e))


class CustomerSummaryTab(QWidget):
    def __init__(self):
        super().__init__()
        layout = QVBoxLayout(self)
        hdr = QLabel("Musteri Bazinda Ciro Ozeti")
        hdr.setStyleSheet("font-size:14px;font-weight:bold;margin-bottom:4px;")
        layout.addWidget(hdr)

        self.table = QTableWidget()
        self.table.setColumnCount(5)
        self.table.setHorizontalHeaderLabels(
            ["Musteri","Plaka","Satis Sayisi","Toplam Ciro","Toplam Kar"])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        layout.addWidget(self.table)

    def load(self):
        try:
            data = api.get_customer_summary()
            self.table.setRowCount(0)
            for row, c in enumerate(data):
                self.table.insertRow(row)
                profit = c.get("total_profit",0)
                cells = [
                    QTableWidgetItem(c["name"]),
                    _center(c.get("car_plate","-")),
                    _center(str(c["sale_count"])),
                    _right(f"{c['total_revenue']:,.2f} TL"),
                    _right(f"{profit:,.2f} TL"),
                ]
                for col, item in enumerate(cells):
                    self.table.setItem(row,col,item)
                color = "#a6e3a1" if profit>=0 else "#f38ba8"
                self.table.item(row,4).setForeground(QBrush(QColor(color)))
        except APIError as e:
            QMessageBox.critical(self,"Hata",str(e))


class ReportsModule(BaseModule):
    def setup_ui(self):
        layout = QVBoxLayout(self)
        title = QLabel("Raporlar & Analiz")
        title.setStyleSheet("font-size:18px;font-weight:bold;margin-bottom:4px;")
        layout.addWidget(title)

        self.tabs = QTabWidget()

        self.tab_daily    = PeriodReportTab("Gunluk",  api.get_daily_report)
        self.tab_weekly   = PeriodReportTab("Haftalik", api.get_weekly_report)
        self.tab_monthly  = PeriodReportTab("Aylik",   api.get_monthly_report)
        self.tab_products = TopProductsTab()
        self.tab_suppliers= SupplierSummaryTab()
        self.tab_customers= CustomerSummaryTab()

        self.tabs.addTab(self.tab_daily,     "Gunluk")
        self.tabs.addTab(self.tab_weekly,    "Haftalik")
        self.tabs.addTab(self.tab_monthly,   "Aylik")
        self.tabs.addTab(self.tab_products,  "Urun Analizi")
        self.tabs.addTab(self.tab_suppliers, "Tedarkiciler")
        self.tabs.addTab(self.tab_customers, "Musteriler")

        layout.addWidget(self.tabs)

    def refresh_data(self):
        for tab in [self.tab_daily, self.tab_weekly, self.tab_monthly,
                    self.tab_products, self.tab_suppliers, self.tab_customers]:
            tab.load()